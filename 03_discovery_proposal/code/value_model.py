"""D3 value model: buyer-supplied inputs (from the role-play), scenarios, payback, 3-year TCO vs HelixAsset and TrueSense, discount rules."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
BLUE = Font(name="Arial", size=10, color="0000FF"); BLK = Font(name="Arial", size=10); BOLD = Font(name="Arial", size=10, bold=True)
GRN = Font(name="Arial", size=10, color="008000"); HDR = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HF = PatternFill("solid", fgColor="1F3A5F"); YEL = PatternFill("solid", fgColor="FFFF00"); BUY = PatternFill("solid", fgColor="D9F2E6")
SM = Font(name="Arial", size=9, color="52514E")
wb = Workbook()
def sh(name, widths):
    ws = wb.create_sheet(name)
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    return ws
def put(ws, ref, v, f=BLK, fmt=None, fill=None, wrap=False):
    c = ws[ref]; c.value = v; c.font = f
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    if wrap: c.alignment = Alignment(wrap_text=True, vertical="top")
    return c
def hdr(ws, row, vals, start=1):
    for j, v in enumerate(vals, start):
        c = ws.cell(row, j, v); c.font = HDR; c.fill = HF; c.alignment = Alignment(wrap_text=True, vertical="top")
def title(ws, t, s):
    put(ws, "A1", t, Font(name="Arial", size=13, bold=True, color="1F3A5F")); put(ws, "A2", s, Font(name="Arial", size=9, italic=True, color="52514E"))
CR = "#,##0.00;(#,##0.00);-"; LK = "#,##0.0;(#,##0.0);-"; PCT = "0%"; INT = "#,##0"
# ---------------- README
ws = wb.active; ws.title = "README"; ws.column_dimensions["A"].width = 24; ws.column_dimensions["B"].width = 110
for i, (a, b) in enumerate([("D3 VALUE MODEL: RANGANATH PRECISION COMPONENTS (CHAKAN PLANT 2)", ""),
    ("Company", "Strata Sense and Ranganath are fictional; synthetic data (MBA capstone)."),
    ("Colour code", "Green fill = number supplied by the buyer in discovery (who said it is in the Source column). Blue text = input. Yellow fill = Strata assumption to validate in the pilot. Black = formula."),
    ("Money", "₹ crore unless marked lakh (1 crore = 100 lakh)."),
    ("Scenarios", "Conservative = what Kavita's 12-month payback rule is tested on (no penalties, no energy, shortest outages). Base and Upside shown for context."),
    ("Options", "Good = press line (22 presses + ~140 hydraulic pumps/motors = 162 assets). Better = Good + 60 compressors + ~200 critical pumps, chillers and fans (422). Best = every non-spindle asset (670); the 480 CNC spindles stay with the Helix pilot."),
    ("Sheets", "Inputs · Scope_Price · Value · Payback · TCO_3yr · Discount_Rules · Checks")], 1):
    ws.cell(i, 1, a).font = BOLD if i == 1 else BLK; ws.cell(i, 2, b).font = BLK; ws.cell(i, 2).alignment = Alignment(wrap_text=True)
# ---------------- Inputs
I = sh("Inputs", [52, 13, 13, 13, 58]); title(I, "Inputs", "Green = buyer-supplied (role-play, Week 3). Yellow = assumption. Column B-D = Conservative / Base / Upside where they differ.")
hdr(I, 4, ["Input", "Conservative", "Base", "Upside", "Source"])
rows = [
 # key, label, c, b, u, fmt, src, kind
 ("fail", "Major press failures per year", 3, 3, 3, "0", "Suresh (plant head): 3 in FY26", "buyer"),
 ("dur", "Hours down per major press failure", 36, 48, 60, "0", "Suresh: 36–60 hours each (Nov event 52 h)", "buyer"),
 ("cph", "Press-line contribution lost per hour (₹ lakh)", 5.5, 5.5, 5.5, "0.0", "Kavita (CFO): ₹5.5 lakh/h", "buyer"),
 ("cmh", "Machining-line contribution lost per hour (₹ lakh)", 2.4, 2.4, 2.4, "0.0", "Kavita: ₹2.4 lakh/h", "buyer"),
 ("pen", "OEM penalty per stoppage event (₹ lakh)", 120, 120, 120, "0", "Kavita: ₹1.2 cr in Nov (₹8 lakh/h × ~15 h)", "buyer"),
 ("penev", "OEM stoppage events per year", 1, 1, 1, "0", "Suresh/Kavita: one in FY26", "buyer"),
 ("maint", "Plant maintenance spend per year (₹ cr)", 6.5, 6.5, 6.5, "0.0", "Suresh: ₹6.5 cr", "buyer"),
 ("react", "Share of maintenance that is reactive", 0.55, 0.55, 0.55, PCT, "Suresh: ~55%", "buyer"),
 ("spares", "Plant spares inventory (₹ cr)", 5.2, 5.2, 5.2, "0.0", "Kavita: ₹5.2 cr (group ₹38 cr)", "buyer"),
 ("power", "Group power bill (₹ cr/yr)", 62, 62, 62, "0", "Kavita: ₹62 cr", "buyer"),
 ("comp", "Compressor share of plant electricity", 0.18, 0.18, 0.18, PCT, "Kavita: ~18%", "buyer"),
 ("payrule", "Payback rule (months, conservative case)", 12, 12, 12, "0", "Kavita: under 12 months", "buyer"),
 ("limit", "CFO's own approval limit (₹ cr/yr opex)", 1.5, 1.5, 1.5, "0.0", "Kavita: ₹1.5 cr; above → committee (15 Jul / 15 Oct)", "buyer"),
 ("detect", "Share of major press failures caught early", 0.33, 0.50, 0.67, PCT, "Assumption; pilot success criterion", "asm"),
 ("avoid", "Share of outage hours avoided when caught early", 0.60, 0.75, 0.85, PCT, "Assumption: planned repair in a scheduled window still takes time", "asm"),
 ("pencount", "Share of OEM penalty exposure counted", 0.0, 0.50, 0.67, PCT, "Conservative counts none (CFO)", "asm"),
 ("mhrs", "Machining hours saved per year from utility failures (Better/Best)", 0, 20, 40, "0", "Assumption; replace with Suresh's paper logs", "asm"),
 ("mred", "Cut in reactive maintenance cost on monitored assets", 0.08, 0.12, 0.16, PCT, "Assumption", "asm"),
 ("ereduce", "Compressor energy saved (Better/Best)", 0.0, 0.04, 0.07, PCT, "Assumption; CFO discounts heavily, so conservative = 0", "asm"),
 ("sred", "Spares reduction on monitored assets (one-off cash)", 0.05, 0.08, 0.10, PCT, "Assumption", "asm"),
 ("plants", "Plants in group (for plant power share)", 9, 9, 9, "0", "Kavita: 9 plants; plant power = group ÷ 9 (assumption)", "buyer"),
]
ref = {}
for k, (key, lab, c, b, u, fmt, src, kind) in enumerate(rows):
    r = 5 + k; ref[key] = r
    put(I, f"A{r}", lab)
    for col, v in zip("BCD", (c, b, u)):
        put(I, f"{col}{r}", v, BLUE, fmt, BUY if kind == "buyer" else YEL)
    put(I, f"E{r}", src, SM)
put(I, f"A{5+len(rows)+1}", "Asset counts, Chakan P2 (Suresh)", BOLD)
ac = [("spindle", "CNC spindles (in Helix pilot)", 480), ("press", "Hydraulic presses", 22), ("hyd", "Pumps/motors on press hydraulics", 140),
      ("compr", "Compressors and blowers", 60), ("util", "Other pumps, fans, chillers, conveyors", 448)]
for k, (key, lab, v) in enumerate(ac):
    r = 5 + len(rows) + 2 + k; ref[key] = r; put(I, f"A{r}", lab); put(I, f"B{r}", v, BLUE, INT, BUY); put(I, f"E{r}", "Suresh: ~1,150 total", SM)
r = 5 + len(rows) + 2 + len(ac); ref["total"] = r; put(I, f"A{r}", "Total rotating assets", BOLD); put(I, f"B{r}", f"=SUM(B{ref['spindle']}:B{r-1})", BOLD, INT)
I.freeze_panes = "B5"
R = lambda key, col="B": f"Inputs!${col}${ref[key]}"
# ---------------- Scope & price
S = sh("Scope_Price", [46, 14, 14, 14, 54]); title(S, "Scope and price", "List prices from DR5; per-asset pricing; hardware can be bundled as a service to keep the contract opex")
hdr(S, 4, ["Item", "Good: press line", "Better: + compressors & critical utilities", "Best: all non-spindle assets", "Source"])
put(S, "A5", "Critical utilities included in Better (count)"); put(S, "C5", 200, BLUE, INT, YEL); put(S, "E5", "Assumption: most critical of the 448", SM)
put(S, "A6", "Monitored assets"); put(S, "B6", f"={R('press')}+{R('hyd')}", GRN, INT); put(S, "C6", f"=B6+{R('compr')}+C5", BLK, INT); put(S, "D6", f"=B6+{R('compr')}+{R('util')}", BLK, INT)
put(S, "A7", "List subscription per asset (₹/yr)"); [put(S, f"{c}7", 12500, BLUE, INT) for c in "BCD"]; put(S, "E7", "DR5 Strata list price", SM)
put(S, "A8", "Hardware per asset, one-off (₹)"); [put(S, f"{c}8", 11000, BLUE, INT) for c in "BCD"]; put(S, "E8", "DR5: ₹9,000–14,000, avg ~₹11,000", SM)
put(S, "A9", "Assets per gateway / gateway cost (₹ lakh)"); put(S, "B9", 150, BLUE, INT); put(S, "C9", 1.2, BLUE, "0.0"); put(S, "E9", "DR5", SM)
put(S, "A10", "Implementation per plant (₹ lakh)"); put(S, "B10", 6, BLUE, "0.0"); put(S, "E10", "DR5", SM)
put(S, "A11", "Edge server for on-premise buffering (₹ lakh)"); put(S, "B11", 8, BLUE, "0.0", YEL); put(S, "E11", "Assumption; required by RPC CISO rules (Kavita)", SM)
put(S, "A13", "Annual subscription = ARR (₹ cr)", BOLD)
for c in "BCD": put(S, f"{c}13", f"={c}6*{c}7/1e7", BOLD, CR)
put(S, "A14", "Hardware + gateways + implementation + edge, one-off (₹ cr)")
for c in "BCD": put(S, f"{c}14", f"=({c}6*{c}8+ROUNDUP({c}6/$B$9,0)*$C$9*1e5+$B$10*1e5+$B$11*1e5)/1e7", BLK, CR)
put(S, "A15", "Year-1 cost, hardware paid upfront (₹ cr)"); [put(S, f"{c}15", f"={c}13+{c}14", BLK, CR) for c in "BCD"]
put(S, "A16", "Hardware-as-a-service: one-offs spread over 36 months at 10% (₹ cr/yr)"); [put(S, f"{c}16", f"=PMT(10%/12,36,-{c}14)*12", BLK, CR) for c in "BCD"]
put(S, "A17", "Annual fee with hardware as a service (₹ cr/yr)", BOLD); [put(S, f"{c}17", f"={c}13+{c}16", BOLD, CR) for c in "BCD"]
put(S, "A18", "Within CFO's own ₹1.5 cr limit?"); [put(S, f"{c}18", f'=IF({c}17<={R("limit")},"Yes","No: committee")', BOLD) for c in "BCD"]
put(S, "A20", "Paid pilot (press line, 4 months), fee (₹ cr)"); put(S, "B20", 0.25, BLUE, CR, YEL); put(S, "E20", "Covers hardware and install; credited in full to year 1 on conversion", SM)
# ---------------- Value
V = sh("Value", [48, 13, 13, 13, 13, 13, 13, 13, 13, 13]); title(V, "Annual value by driver (₹ cr)", "Formulas use the Inputs sheet; each driver switches on only for the options that cover those assets")
hdr(V, 4, ["Driver", "Good C", "Good B", "Good U", "Better C", "Better B", "Better U", "Best C", "Best B", "Best U"])
cols = {"Good": "BCD", "Better": "EFG", "Best": "HIJ"}; sc = "BCD"
mshare = {"Good": 0.25, "Better": 0.45, "Best": 0.60}; sshare = {"Good": 0.20, "Better": 0.35, "Best": 0.50}
put(V, "A5", "Press downtime avoided")
put(V, "A6", "OEM penalties avoided")
put(V, "A7", "Machining downtime avoided (utility failures)")
put(V, "A8", "Reactive maintenance saved")
put(V, "A9", "Compressor energy saved")
put(V, "A10", "Total annual value", BOLD)
put(V, "A12", "One-off cash from lower spares inventory")
put(V, "A14", "Share of maintenance spend on monitored assets"); put(V, "A15", "Share of spares for monitored assets")
for opt, cc in cols.items():
    for i, col in enumerate(cc):
        s = sc[i]
        put(V, f"{col}5", f"={R('fail',s)}*{R('dur',s)}*{R('avoid',s)}*{R('detect',s)}*{R('cph',s)}/100", BLK, CR)
        put(V, f"{col}6", f"={R('penev',s)}*{R('pen',s)}*{R('pencount',s)}*{R('detect',s)}/100", BLK, CR)
        put(V, f"{col}7", "=0" if opt == "Good" else f"={R('mhrs',s)}*{R('cmh',s)}/100*{1 if opt=='Better' else 1.3}", BLK, CR)
        put(V, f"{col}8", f"={R('maint',s)}*{R('react',s)}*{col}14*{R('mred',s)}", BLK, CR)
        put(V, f"{col}9", "=0" if opt == "Good" else f"={R('power',s)}/{R('plants',s)}*{R('comp',s)}*{R('ereduce',s)}", BLK, CR)
        put(V, f"{col}10", f"=SUM({col}5:{col}9)", BOLD, CR)
        put(V, f"{col}12", f"={R('spares',s)}*{col}15*{R('sred',s)}", BLK, CR)
        put(V, f"{col}14", mshare[opt], BLUE, PCT, YEL); put(V, f"{col}15", sshare[opt], BLUE, PCT, YEL)
put(V, "A17", "Best adds 30% more machining-hours saved than Better (wider utility coverage): assumption.", SM)
# ---------------- Payback
P = sh("Payback", [46, 13, 13, 13, 13, 13, 13, 13, 13, 13]); title(P, "Payback and ROI", "Payback months = annual fee (hardware as a service) ÷ annual value × 12. Tested on Conservative.")
hdr(P, 4, ["Measure", "Good C", "Good B", "Good U", "Better C", "Better B", "Better U", "Best C", "Best B", "Best U"])
feecol = {"Good": "B", "Better": "C", "Best": "D"}
put(P, "A5", "Annual value (₹ cr)"); put(P, "A6", "Annual fee incl. hardware as a service (₹ cr)"); put(P, "A7", "Net annual benefit (₹ cr)")
put(P, "A8", "Payback (months)", BOLD); put(P, "A9", "Value ÷ fee (x)"); put(P, "A10", "Meets CFO 12-month rule? (Conservative)", BOLD)
for opt, cc in cols.items():
    for i, col in enumerate(cc):
        put(P, f"{col}5", f"=Value!{col}10", GRN, CR); put(P, f"{col}6", f"=Scope_Price!{feecol[opt]}17", GRN, CR)
        put(P, f"{col}7", f"={col}5-{col}6", BLK, CR); put(P, f"{col}8", f'=IF({col}5>0,{col}6/{col}5*12,"n/a")', BOLD, "0.0")
        put(P, f"{col}9", f'=IF({col}6>0,{col}5/{col}6,0)', BLK, "0.0x")
    put(P, f"{cc[0]}10", f'=IF({cc[0]}8<={R("payrule")},"Yes","No")', BOLD)
# ---------------- TCO
T = sh("TCO_3yr", [52, 14, 14, 14, 56]); title(T, "Three-year total cost and value-adjusted comparison (Good scope: press line)", "DR5 list prices. Helix year 1 is free only with a 3-year MES renewal (Kavita).")
hdr(T, 4, ["Item", "Strata Cortex", "HelixAsset", "TrueSense AI", "Source / note"])
put(T, "A5", "Assets in scope"); [put(T, f"{c}5", "=Scope_Price!B6", GRN, INT) for c in "BCD"]
put(T, "A6", "Subscription per asset per year (₹)"); put(T, "B6", "=Scope_Price!B7", GRN, INT); put(T, "C6", 9500, BLUE, INT); put(T, "D6", 7800, BLUE, INT); put(T, "E6", "DR5 list prices", SM)
put(T, "A7", "Years charged out of 3"); put(T, "B7", 3, BLUE); put(T, "C7", 2, BLUE); put(T, "D7", 3, BLUE); put(T, "E7", "Helix: year 1 free with 3-year MES lock-in", SM)
put(T, "A8", "Sensor hardware per asset (₹)"); put(T, "B8", "=Scope_Price!B8", GRN, INT); put(T, "C8", 13000, BLUE, INT); put(T, "D8", 8000, BLUE, INT); put(T, "E8", "DR5: presses are not wired to Helix PLCs and have no online sensors today (Suresh)", SM)
put(T, "A9", "Other one-offs (gateways, implementation, edge) (₹ cr)"); put(T, "B9", "=Scope_Price!B14-Scope_Price!B6*Scope_Price!B8/1e7", GRN, CR); put(T, "C9", 0.08, BLUE, CR, YEL); put(T, "D9", 0.10, BLUE, CR, YEL); put(T, "E9", "Rival integration/edge costs: assumption", SM)
put(T, "A10", "Three-year cost (₹ cr)", BOLD); [put(T, f"{c}10", f"={c}5*{c}6*{c}7/1e7+{c}5*{c}8/1e7+{c}9", BOLD, CR) for c in "BCD"]
put(T, "A11", "Share of press failures caught early (base)"); put(T, "B11", f"={R('detect','C')}", GRN, PCT); put(T, "C11", 0.20, BLUE, PCT, YEL); put(T, "D11", 0.30, BLUE, PCT, YEL)
put(T, "E11", "Helix sees PLC tags only; TrueSense weaker on hydraulic/low-speed assets (DR5). Assumptions to prove head-to-head in the pilot.", SM)
put(T, "A12", "Three-year value, base (₹ cr)"); [put(T, f"{c}12", f"=3*(Value!$C$5*{c}11/$B$11+Value!$C$6*{c}11/$B$11+Value!$C$8)", BLK, CR) for c in "BCD"]
put(T, "A13", "Three-year net value (₹ cr)", BOLD); [put(T, f"{c}13", f"={c}12-{c}10", BOLD, CR) for c in "BCD"]
put(T, "A14", "Strata cost premium vs rival (₹ cr, 3 years)"); put(T, "C14", "=B10-C10", BLK, CR); put(T, "D14", "=B10-D10", BLK, CR)
put(T, "A15", "Strata net-value advantage vs rival (₹ cr, 3 years)"); put(T, "C15", "=B13-C13", BOLD, CR); put(T, "D15", "=B13-D13", BOLD, CR)
put(T, "A17", "Reading: Strata costs more over three years; it only wins if it catches materially more press failures. The pilot must prove that, so the detection rate is the pilot's first success criterion. Not costed: the value of avoiding a 3-year MES lock-in.", SM)
T["A17"].alignment = Alignment(wrap_text=True); T.row_dimensions[17].height = 40
# ---------------- Discount rules
D = sh("Discount_Rules", [46, 14, 60]); title(D, "Discount rules: every point is traded for something", "List ₹12,500 per asset per year. FY26 average discount was 19% and bought no extra wins (D1).")
hdr(D, 4, ["Give (discount off subscription)", "Max %", "Get (what RPC commits)"])
dr = [("Three-year term", 0.08, "Contracted 3-year subscription with annual payment"), ("Annual payment in advance", 0.03, "12 months paid upfront"),
      ("Second-site commitment", 0.05, "Signed option to extend to a Chennai forging plant within 12 months of conversion"),
      ("Reference and case study", 0.02, "Named reference, site visit and a published case study after 6 months")]
for k, (a, b, c) in enumerate(dr):
    r = 5 + k; put(D, f"A{r}", a); put(D, f"B{r}", b, BLUE, PCT); put(D, f"C{r}", c)
put(D, "A9", "Maximum total discount", BOLD); put(D, "B9", "=MIN(0.15,SUM(B5:B8))", BOLD, PCT); put(D, "C9", "Hard cap 15% (brief target: average discount ≤ 14%)", SM)
put(D, "A10", "Floor price per asset (₹/yr)", BOLD); put(D, "B10", "=Scope_Price!B7*(1-B9)", BOLD, INT)
put(D, "A12", "Never traded", BOLD); put(D, "C12", "Hardware below cost · a free pilot without a signed conversion price · outcome-linked fees without a measured baseline · discounts for 'matching Helix year-1 free'", wrap=True)
put(D, "A14", "Outcome-linked option (no discount)"); put(D, "C14", "15% of the annual subscription at risk, paid only if agreed avoided-downtime hours are met; measured from Suresh's paper press log as baseline.", wrap=True)
put(D, "A16", "Negotiation test: Neha asks 30% off"); put(D, "B16", 0.30, BLUE, PCT); put(D, "C16", '=IF(B16>B9,"Above the 15% cap: counter with give-gets; walk away below floor","Within rules")', BOLD)
# ---------------- Checks
C = sh("Checks", [60, 16, 12]); title(C, "Checks", "")
hdr(C, 4, ["Check", "Value", "OK?"])
ck = [("Asset total = ~1,150 (Suresh)", f"={R('total')}", '=IF(ABS(B5-1150)<=10,"OK","CHECK")'),
      ("Better year-1 fee within CFO limit", "=Scope_Price!C17", f'=IF(B6<={R("limit")},"OK","COMMITTEE")'),
      ("Good payback, conservative (months)", "=Payback!B8", f'=IF(B7<={R("payrule")},"OK","FAILS RULE")'),
      ("Better payback, conservative (months)", "=Payback!E8", f'=IF(B8<={R("payrule")},"OK","FAILS RULE")'),
      ("Max discount ≤ 15%", "=Discount_Rules!B9", '=IF(B9<=0.15,"OK","OVER")'),
      ("Nov OEM penalty reconciles (₹8 lakh × 15 h = ₹120 lakh)", f"={R('pen')}", '=IF(B10=8*15,"OK","CHECK")')]
for k, (a, b, c) in enumerate(ck):
    r = 5 + k; put(C, f"A{r}", a); put(C, f"B{r}", b, BLK, "0.00"); put(C, f"C{r}", c, BOLD)
wb.save("out/D3_4_Value_Model.xlsx"); print("saved")
