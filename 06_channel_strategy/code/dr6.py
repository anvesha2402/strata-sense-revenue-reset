import pandas as pd, numpy as np
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
o = pd.read_pickle("/home/claude/strata/d1/out/opps_analysis.pkl"); p = o[o.partner_id.notna()].copy()
w = p[p.is_won]; q = p[p.qualified & p.stage.isin(["Closed Won", "Closed Lost"])]; op = p[~p.is_closed]
def arr(pid, fy, src): 
    x = w[(w.partner_id == pid) & (w.fy == fy) & (w.partner_sourced_adj == src)]; return round(x.first_year_value_lakh.sum() / 100, 2)
base = [("P01", "Gulf Automation Services LLC", "System integrator", "Middle East (UAE)", "2021-06", 1, "Orion Controls", "Inactive"),
 ("P02", "Deccan Industrial Systems Pvt Ltd", "System integrator", "India-West", "2019-03", 6, "Orion Controls, Vektor Systems", "Active"),
 ("P03", "Coromandel Controls", "Automation distributor", "India-South", "2020-01", 1, "Kaizen Automation", "Inactive"),
 ("P04", "Northern Drives & Automation", "Automation distributor", "India-North", "2020-08", 0, "Orion Controls", "Inactive"),
 ("P05", "Sahyadri Automation Pvt Ltd", "System integrator", "India-West", "2019-11", 5, "Helix Automation (authorised integrator), Kaizen", "Active"),
 ("P06", "Najd Engineering Co.", "Value-added reseller", "Middle East (Saudi Arabia)", "2022-02", 1, "Vektor Systems", "Inactive"),
 ("P07", "Mekong Industrial Tech", "System integrator", "SEA (Vietnam)", "2024-05", 0, "Orion Controls", "Signed, no deals"),
 ("P08", "Hosur Machine Builders", "Machine builder (OEM)", "India-South", "2023-01", 0, "Own machines", "Signed, no deals"),
 ("P09", "Vindhya Reliability Services", "Managed-service provider", "India-North / East", "2020-06", 7, "Vibration services; Vektor Systems", "Active"),
 ("P10", "Penang Plant Solutions Sdn Bhd", "System integrator", "SEA (Malaysia)", "2025-01", 0, "Orion Controls", "Signed, no deals"),
 ("P11", "Konkan Process Automation Ltd", "System integrator", "India-West / South", "2018-09", 4, "Helix Automation (gold partner), Orion", "Active"),
 ("P12", "Siam Reliability Co.", "Managed-service provider", "SEA (Thailand)", "2023-07", 2, "Independent", "Occasional"),
 ("P13", "Al Fajr Industrial Services", "System integrator", "Middle East (Oman / UAE)", "2022-10", 1, "Kaizen Automation", "Occasional"),
 ("P14", "Ganga Electricals", "Automation distributor", "India-North", "2021-04", 0, "Orion Controls", "Inactive")]
rows = []
for pid, nm, typ, reg, sd, cert, lines, st in base:
    qq = q[q.partner_id == pid]; oo = op[op.partner_id == pid]
    rows.append(dict(partner_id=pid, partner_name=nm, type=typ, region=reg, signed=sd, status=st, certified_engineers=cert, other_vendor_lines=lines,
                     sourced_arr_fy25_cr=arr(pid, 2025, True), sourced_arr_fy26_cr=arr(pid, 2026, True), influenced_arr_fy26_cr=arr(pid, 2026, False),
                     qualified_deals_fy24_26=len(qq), win_rate_fy24_26=round(qq.is_won.mean(), 3) if len(qq) else None,
                     open_pipeline_cr=round(oo.first_year_value_lakh.sum() / 100, 2), open_pipeline_helix_incumbent_share=round(oo.helix_inc.mean(), 2) if len(oo) else None))
signed = pd.DataFrame(rows)
pros = pd.DataFrame([
 ("PR01", "Java Integrasi Teknik", "System integrator", "SEA (Indonesia)", 60, 180, "Cement, pulp & paper, palm oil", "Orion Controls", "No", 140, "High", 6),
 ("PR02", "Chao Phraya Automation Co.", "System integrator", "SEA (Thailand)", 45, 150, "Auto components, food", "Vektor Systems", "No", 120, "High", 6),
 ("PR03", "Saigon Reliability JSC", "Managed-service provider", "SEA (Vietnam)", 25, 40, "Electronics, textiles", "Independent", "No", 90, "Medium", 9),
 ("PR04", "Luzon Plant Services Inc.", "Managed-service provider", "SEA (Philippines)", 20, 35, "Power, food, cement", "Independent", "No", 60, "Medium", 9),
 ("PR05", "Malacca Automation Sdn Bhd", "Automation distributor", "SEA (Malaysia)", 30, 220, "Mixed", "Helix Automation (distributor)", "Yes", 110, "Medium", 6),
 ("PR06", "Merlion Industrial Pte Ltd", "Value-added reseller", "SEA regional (Singapore HQ)", 35, 260, "Semiconductors, pharma, chemicals", "Orion, Vektor", "No", 130, "High", 9),
 ("PR07", "Arabian Asset Care Co.", "Managed-service provider", "Middle East (Saudi Arabia)", 40, 120, "Oil & gas, utilities, cement", "Independent", "No", 110, "High", 9),
 ("PR08", "Emirates Crest Engineering LLC", "System integrator", "Middle East (UAE)", 30, 140, "Utilities, aluminium, water", "Helix Automation (integrator)", "Yes", 80, "Medium", 6),
 ("PR09", "Tapti Press Machines Ltd", "Machine builder (OEM embed)", "India (national)", 15, 400, "Hydraulic press maker (installed base ~2,000 presses)", "Own machines", "No", 300, "High", 12),
 ("PR10", "Shivalik Compressors Ltd", "Machine builder (OEM embed)", "India (national)", 12, 650, "Industrial compressors", "Own machines", "No", 450, "Medium", 12),
 ("PR11", "Narmada Reliability Services", "Managed-service provider", "India-East", 18, 30, "Steel, power", "Independent", "No", 150, "Medium", 6),
 ("PR12", "Nilgiri Drives Pvt Ltd", "Automation distributor", "India-South", 22, 300, "Mixed discrete", "Helix Automation (distributor)", "Yes", 200, "Low", 9)],
 columns=["prospect_id", "name", "type", "region", "engineers", "annual_revenue_cr", "vertical_focus", "vendor_lines", "helix_affiliated", "overlap_with_strata_targets", "interest", "est_months_to_first_deal"])
readme = pd.DataFrame([("DR6 PARTNER FILE", ""), ("Synthetic", "Fictional partners; synthetic data for an MBA capstone. ARR and pipeline figures computed from DR1 (cleaned in D1) using partner_id."),
    ("Sheets", "signed_partners: the 14 signed partners. prospective_partners: 12 profiles gathered by the CRO office."),
    ("Definitions", "Sourced = partner brought the deal (incl. 36 deals re-tagged in D1). Influenced = partner helped on a direct deal. overlap_with_strata_targets = number of Strata's 3,100 target accounts the prospect already serves (partner-reported, unverified).")], columns=["Item", "Description"])
pth = "dr/DR6_partner_file.xlsx"
with pd.ExcelWriter(pth) as wr:
    readme.to_excel(wr, sheet_name="README", index=False, header=False); signed.to_excel(wr, sheet_name="signed_partners", index=False); pros.to_excel(wr, sheet_name="prospective_partners", index=False)
wb = load_workbook(pth)
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row: c.font = Font(name="Arial", size=10, bold=c.font.bold); c.alignment = Alignment(wrap_text=True, vertical="top")
    if ws.title != "README":
        for c in ws[1]: c.font = Font(name="Arial", size=10, bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3A5F")
        ws.freeze_panes = "C2"
    for i, col in enumerate(ws.columns, 1): ws.column_dimensions[get_column_letter(i)].width = 90 if col[0].value == "Description" else (30 if i in (2, 8) else 15)
wb.save(pth); signed.to_pickle("dr/signed.pkl"); pros.to_pickle("dr/pros.pkl")
print(signed[["partner_id", "status", "sourced_arr_fy26_cr", "influenced_arr_fy26_cr", "qualified_deals_fy24_26", "win_rate_fy24_26", "open_pipeline_cr"]].to_string())
