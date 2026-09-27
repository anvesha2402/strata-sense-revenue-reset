import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
sel = pd.read_pickle("dr/dr4_sel.pkl").set_index("opp_id")
S = {  # opp_id: (interviewee, date, summary, quote)
"OPP-110215": ("Plant maintenance head", "2026-02-18", "Strata ran a strong demo and a proposal with a business case, but the CIO had already agreed a Helix MES renewal that included HelixAsset for the first year at no charge. Maintenance preferred Strata's diagnostics on presses; IT wanted one vendor on the OT network. Strata offered 25% off late in the cycle; it made no difference.", "\"Your analytics were better. But 'free and already inside our control system' ends the argument with IT.\""),
"OPP-106456": ("Operations director", "2025-01-22", "Evaluation stalled when Helix proposed including HelixAsset in the automation upgrade already budgeted for the year. Strata's contact was the reliability engineer only; the director first met Strata in the final week. Strata's discount was matched by a Helix 'loyalty' offer.", "\"We never heard from your people at my level until procurement asked for a final price.\""),
"OPP-110102": ("Reliability engineer", "2026-03-05", "Lost before a proposal: the plant decided to 'try HelixAsset first' because it came with the MES licence. The engineer thought Helix would not cover compressors or chillers but could not get time with the plant head to argue it.", "\"I told them Helix sees only what the PLC sees. Nobody above me was listening.\""),
"OPP-107405": ("Station manager", "2025-02-10", "Helix is the plant's automation vendor and presented HelixAsset as part of a five-year service contract. Strata had two contacts and a business case, but the station's finance team treated condition monitoring as a service-contract line item rather than a separate investment.", "\"Finance saw it as already paid for inside the Helix contract.\""),
"OPP-111858": ("Head of engineering", "2026-01-28", "Helix was not the incumbent but its regional integrator was building a new line and bundled HelixAsset with the installation. Strata was not in the integrator's design; by the time it engaged, the architecture was fixed.", "\"The integrator decides a lot here. You were not in the room when the line was designed.\""),
"OPP-111374": ("Utilities manager", "2026-02-02", "The account already used third-party wireless sensors; Helix offered its analytics module bundled with a DCS upgrade. Strata's single contact left the company mid-cycle and the deal lapsed to the bundle.", "\"Our champion left. After that it was Helix by default.\""),
"OPP-109944": ("Plant head", "2026-01-15", "TrueSense ran a free 60-day proof of concept on existing online sensors. Strata's proposal required new sensors and a paid pilot; with one contact and no business case, the plant saw no reason to pay. In the end nobody was chosen.", "\"Why pay you to prove something the other vendor proved for free on our own sensors?\""),
"OPP-107022": ("Maintenance manager", "2024-11-20", "The site had online wired sensors on most critical assets. TrueSense connected to the historian in two weeks with no hardware. Strata's rep argued sensor quality but could not show where the existing data fell short.", "\"They used what we already had. You wanted to replace it.\""),
"OPP-109968": ("Procurement manager", "2026-02-24", "Procurement compared per-asset prices: TrueSense about 35% lower. The plant head preferred Strata for the forging presses but did not join the final meetings; the rep dealt only with procurement at the end.", "\"On paper the price gap was big, and nobody gave me a reason to pay it.\""),
"OPP-107338": ("Warehouse operations head", "2024-12-12", "Small site with wireless sensors already fitted by another vendor. TrueSense's software-only offer was cheaper and quick. Strata's proposal had no business case.", "\"We did not need new hardware. Your offer started with hardware.\""),
"OPP-109453": ("Facilities manager", "2025-12-08", "Commercial real-estate site with simple fans and pumps. BMI's low-cost sensors and dashboards were 'good enough'. Strata's value case was built for heavy industry.", "\"For fans and pumps we needed trends, not a PhD.\""),
"OPP-110581": ("Plant engineer", "2026-01-19", "BMI's regional partner offered sensors and dashboards at about half Strata's price with local support. The steel plant's critical assets were not covered in BMI's scope; the plant plans to revisit them later.", "\"Cheap and local won phase one. Phase two is still open.\""),
"OPP-109053": ("Plant head", "2025-11-26", "Evaluation stopped when the MES renewal with Helix was delayed; management froze all related decisions. Strata's single contact had no budget authority and no business case in the proposal.", "\"Nobody was going to decide this before the Helix renewal was settled.\""),
"OPP-108595": ("Utilities head", "2025-01-30", "Budget frozen after a poor quarter. TrueSense and Strata both paused. Strata never met finance, so there was no one to argue for the savings.", "\"The budget froze and nobody in finance knew what the savings were.\""),
"OPP-107099": ("Mill manager", "2024-12-18", "Good business case and two contacts, but the project was deferred when the parent group started a group-wide digital programme. Strata had not mapped the group IT decision process.", "\"Head office took over digital decisions in the middle of your proposal.\""),
"OPP-105266": ("Operations manager", "2024-08-14", "Proposal included a business case built on vendor averages, not the plant's own figures. Finance did not believe the savings and the project was dropped.", "\"The numbers were yours, not ours. Finance ignored them.\""),
"OPP-111034": ("Maintenance head", "2026-02-11", "TrueSense's proof of concept struggled on slow-speed kiln drives and mill gearboxes where existing sensors were poor. Strata's paid pilot caught a gearbox fault three weeks early, with finance and the plant head in the review.", "\"Their software was fine where our sensors were good. Where they weren't, you found the fault.\""),
"OPP-106422": ("Plant head", "2024-10-03", "The plant had only portable monitoring, so TrueSense's no-hardware pitch did not apply. Strata's business case used the plant's own breakdown log.", "\"You used our breakdown log. That is what finance signed.\""),
"OPP-108300": ("Plant head", "2024-12-05", "Helix was incumbent, but HelixAsset did not cover the hydraulic systems on the press line. Strata scoped only those assets as a complement and did not compete for the PLC-connected machines.", "\"You didn't try to replace Helix. You covered what Helix couldn't see.\""),
"OPP-110096": ("Engineering manager", "2026-01-12", "Helix pitched its bundle, but the plant ran mixed automation (Orion and Vektor) so HelixAsset would have needed new Helix sensors everywhere. Strata's all-asset coverage was cheaper in total.", "\"Helix was only 'free' on Helix machines. Most of ours aren't.\""),
"OPP-108355": ("CFO", "2025-02-27", "BMI was cheaper, but Strata met the CFO, plant head and IT, and showed kiln-drive failures BMI could not diagnose. The CFO accepted a higher price for a lower-risk result.", "\"Cheap monitoring that misses the big failure is expensive.\""),
"OPP-108912": ("Operations director", "2025-12-15", "Partner-introduced deal in Malaysia. Strata's partner handled installation and local support; the business case used the plant's downtime cost. Four stakeholders involved from week two.", "\"The partner made it easy locally; your team made the numbers real.\""),
"OPP-111521": ("Finance controller", "2026-03-10", "Strata built the business case with finance using the plant's own downtime cost and ran a 4-month paid pilot with written success criteria. No competitor made it to the final round.", "\"The pilot had pass/fail criteria we wrote together. That made approval simple.\""),
"OPP-110706": ("Plant head", "2026-02-20", "TrueSense was shortlisted, but Strata engaged maintenance, finance and IT early, answered IT's security questions first and offered edge buffering. TrueSense's cloud-only design raised IT concerns.", "\"Your team answered IT before IT had to ask.\""),
}
rows = []
for oid, (who, dt, summ, qt) in S.items():
    r = sel.loc[oid]
    rows.append(dict(interview_id=f"VOB-{len(rows)+1:02d}", opp_id=oid, account_name=r.account_name, segment=r.segment, region=r.region,
                     outcome="Won" if r.won == 1 else "Lost", crm_competitor=r.primary_competitor if isinstance(r.primary_competitor, str) else "(blank)",
                     crm_loss_reason=r.loss_reason if isinstance(r.loss_reason, str) else ("" if r.won == 1 else "(blank)"),
                     interviewee_role=who, interview_date=dt, summary=summ, key_quote=qt))
vob = pd.DataFrame(rows)
notes = pd.DataFrame([
 ("RN-01", "AE, India-West, 4 yrs", "\"We always lose to TrueSense on price. If I could go to 30% I'd win half of them.\""),
 ("RN-02", "AE, India-North, 7 months", "\"Helix accounts are a waste of time; IT won't let a second vendor near the PLCs.\""),
 ("RN-03", "Sales manager, India-South", "\"Discount approvals are fast, so reps lead with discount. Business cases take a week; discounts take an hour.\""),
 ("RN-04", "SE, India-West", "\"TrueSense's POC looks great on good historian data. On hydraulics and slow drives it falls apart. We never show customers that side by side.\""),
 ("RN-05", "AE, SEA", "\"Partners bring me deals I'd never find. But half of them are Helix integrators too.\""),
 ("RN-06", "AE, India-East, 2 yrs", "\"No-decision deals usually die when my one contact moves or budget freezes.\""),
 ("RN-07", "SE, India-North", "\"IT security questionnaires kill three weeks per deal. We should send them before they ask.\""),
 ("RN-08", "AE, Middle East", "\"BMI wins the small sites. I only lose to them when I chase deals I shouldn't.\""),
 ("RN-09", "Sales manager, India-West", "\"Helix is 'free' only on machines wired to Helix PLCs. Most plants are mixed. We don't say that clearly enough.\""),
 ("RN-10", "AE, India-South, 11 months", "\"I don't know what I'm allowed to trade for discount. I just ask my manager.\""),
 ("RN-11", "CSM", "\"Customers who use less than half their licences churn. We see it six months ahead and nobody acts.\""),
 ("RN-12", "AE, India-North, 3 yrs", "\"When finance is in the room, we win. When it's only maintenance, we lose on price or nothing happens.\""),
], columns=["note_id", "author_role", "note"])
readme = pd.DataFrame([("DR4 VOICE OF BUYER", ""), ("Synthetic", "Fictional companies and people; synthetic interview summaries for an MBA capstone. Each summary is linked to a real DR1 opportunity (opp_id)."),
    ("Sheets", "win_loss_interviews: 24 anonymised win/loss interview summaries (FY25–FY26). rep_call_notes: anonymised notes from reps, SEs, managers and a CSM."),
    ("Caution", "Interviews are a small, non-random sample; reps' notes are opinions. Test every claim against DR1.")], columns=["Item", "Description"])
p = "dr/DR4_voice_of_buyer.xlsx"
with pd.ExcelWriter(p) as w:
    readme.to_excel(w, sheet_name="README", index=False, header=False); vob.to_excel(w, sheet_name="win_loss_interviews", index=False); notes.to_excel(w, sheet_name="rep_call_notes", index=False)
wb = load_workbook(p)
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row: c.font = Font(name="Arial", size=10, bold=c.font.bold); c.alignment = Alignment(wrap_text=True, vertical="top")
    if ws.title != "README":
        for c in ws[1]: c.font = Font(name="Arial", size=10, bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3A5F")
        ws.freeze_panes = "A2"
    for i, col in enumerate(ws.columns, 1):
        ws.column_dimensions[get_column_letter(i)].width = {"summary": 70, "key_quote": 50, "note": 80, "Description": 90}.get(col[0].value, 16)
wb.save(p); vob.to_pickle("dr/vob.pkl"); notes.to_pickle("dr/notes.pkl"); print(len(vob), len(notes))
