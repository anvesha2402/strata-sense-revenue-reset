import json, subprocess
from openpyxl import load_workbook
M = json.load(open("out/rowmap.json")); R = M["R"]
def run(edits):
    wb = load_workbook("out/D7_Integrated_Model.xlsx"); a = wb["Assumptions"]
    for ref, val in edits.items(): a[ref] = val
    wb.save("/tmp/claude-0/x.xlsx"); subprocess.run(["python3", "/mnt/skills/public/xlsx/scripts/recalc.py", "/tmp/claude-0/x.xlsx", "90"], capture_output=True)
    d = load_workbook("/tmp/claude-0/x.xlsx", data_only=True)
    g = lambda ws, ref: d[ws][ref].value
    return dict(arr27=g("ARR_bridge", f"O{M['ab']['close']}"), arr29=g("ARR_bridge", f"Q{M['ab']['close']}"), nrr27=g("ARR_bridge", f"O{M['ann']['nrr']}"),
                gm=[round(g("PnL", f"{c}{M['pl']['gmpost']}"), 4) for c in "OPQ"], cacq2=max(g("PnL", f"{c}{M['pl']['cac']}") for c in "DEFGHIJKLMN"), cac28=g("PnL", f"J{M['pl']['cac']}"))
out = {}
out["base"] = run({})
out["supply-10"] = run({f"{c}{R['supply']}": round(v * 0.9) for c, v in zip("CDE", (380, 430, 470))})
out["hw_flat_1.48"] = run({f"{c}{R['hwatt']}": 1.48 for c in "CDE"})
out["attr_29"] = run({f"{c}{R['attr']}": 0.29 for c in "CDE"})
out["vajra_lost"] = run({f"C{R['vaj_p']}": 1.0})
json.dump(out, open("out/extra_runs.json", "w"), indent=1); print(json.dumps(out, indent=1))
