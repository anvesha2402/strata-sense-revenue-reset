"""Writes saved runs (Scenarios, Sensitivity, Cascade base/+3) and the Traceability map into the D7 workbook."""
import json
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L
F = "out/D7_Integrated_Model.xlsx"
wb = load_workbook(F); runs = json.load(open("out/runs.json")); plus3 = json.load(open("out/casc_plus3.json")); M = json.load(open("out/rowmap.json"))
BLK = Font(name="Arial", size=10); BOLD = Font(name="Arial", size=10, bold=True); HDR = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HF = PatternFill("solid", fgColor="1F3A5F"); SM = Font(name="Arial", size=9, color="52514E"); TT = Font(name="Arial", size=13, bold=True, color="1F3A5F")
IT = Font(name="Arial", size=9, italic=True, color="52514E")
def hdr(ws, r, vals):
    for j, v in enumerate(vals, 1):
        c = ws.cell(r, j, v); c.font = HDR; c.fill = HF; c.alignment = Alignment(wrap_text=True, vertical="center")
def fmt_of(k):
    if any(x in k for x in ("Growth", "NRR", "share", "GM", "S&M/ARR")): return "0.0%"
    if "Guardrails" in k or "shortfalls" in k or "headcount" in k: return "0"
    return '#,##0.0;(#,##0.0);"-"'
for n in ("Scenarios", "Sensitivity", "Traceability"):
    if n in wb.sheetnames: del wb[n]
# Scenarios
S = wb.create_sheet("Scenarios", 2); S.column_dimensions["A"].width = 46
for c in "BCDE": S.column_dimensions[c].width = 14
S.column_dimensions["F"].width = 60
S["A1"] = "Scenario comparison (saved runs)"; S["A1"].font = TT
S["A2"] = "Values produced by run_scenarios.py: the live workbook was recalculated with each setting of Assumptions!C4. Change C4 to see any scenario live."; S["A2"].font = IT
hdr(S, 4, ["Metric", "Base", "Downside", "Aggressive", "CFO-capped", "Reading"])
keys = list(runs["res"]["Base"].keys())
notes = {"ARR FY29": "Board path FY29: ₹521.9 cr", "Growth FY27": "Board path 15%; Base misses (defended in memo §4)", "CAC worst qtr": "Q1 FY27 is locked by the D4 forecast in every scenario",
         "CAC worst qtr from Q2 FY27": "Tightest guardrail: Q3 FY27", "Investment FY27": "Envelope ₹15 cr (₹18 cr Aggressive)", "Guardrails failed": "Base: only the Q1 FY27 payback (disclosed)",
         "Investment payback FY29 vs FY26-flat (months)": "Stricter counterfactual: credits only gains beyond FY26 rates"}
for i, k in enumerate(keys):
    r = 5 + i; S.cell(r, 1, k).font = BLK
    for j, s in enumerate(("Base", "Downside", "Aggressive", "CFO-capped")):
        v = runs["res"][s][k]
        try: v = float(v)
        except: pass
        c = S.cell(r, 2 + j, v); c.font = BLK; c.number_format = fmt_of(k)
    S.cell(r, 6, notes.get(k, "")).font = SM
S.freeze_panes = "B5"
# Sensitivity
T = wb.create_sheet("Sensitivity", 3); T.column_dimensions["A"].width = 34
for c in "BCDEFGHIJ": T.column_dimensions[c].width = 13
T["A1"] = "Sensitivity of the Base plan and break-even points (saved runs)"; T["A1"].font = TT
T["A2"] = "Each row flexes one driver with the Assumptions sensitivity shift cells and recalculates the whole model. Reproduce by typing the shift into Assumptions (rows marked SENSITIVITY SHIFTS)."; T["A2"].font = IT
hdr(T, 4, ["Driver and change", "FY29 ARR", "Δ vs Base", "FY29 growth", "FY27 ARR", "FY29 EBITDA", "Worst CAC payback from Q2 FY27", "CAC payback Q4 FY28", "Min GM", "Investment payback FY29 vs FY26-flat"])
base = runs["res"]["Base"]; r = 5
def srow(lab, d):
    global r
    vals = [d["ARR FY29"], d["ARR FY29"] - base["ARR FY29"], d["Growth FY29"], d["ARR FY27"], d["EBITDA FY29"], d["CAC worst qtr from Q2 FY27"], d["CAC Q4 FY28"], d["GM min year"], d["Investment payback FY29 vs FY26-flat (months)"]]
    T.cell(r, 1, lab).font = BOLD if lab == "Base plan" else BLK
    for j, v in enumerate(vals):
        c = T.cell(r, 2 + j, v); c.font = BLK; c.number_format = "0.0%" if j in (2, 7) else '#,##0.0;(#,##0.0);"-"'
    r += 1
srow("Base plan", base)
lab_map = {"Win rate": ("pts", 3), "Sales cycle": ("months", 1), "Discount": ("pts", 3), "Partner productivity": ("%", 30), "Churn": ("pts", 2)}
for drv, (u, v) in lab_map.items():
    for sg in (-1, 1):
        d = runs["sens"][f"{drv}|{sg}"]; srow(f"{drv} {'+' if sg > 0 else '−'}{v}{'%' if u == '%' else ' ' + u}", d)
r += 1; T.cell(r, 1, "BREAK-EVEN: the change at which a payback test fails").font = BOLD; r += 1
hdr(T, r, ["Test", "Win rate (pts vs plan)", "Churn (pts vs plan)", "Plan delivery (share of planned improvement)"]); r += 1
be = runs["be"]
for t in ["CAC payback > 30 months in any quarter from Q2 FY27", "CAC payback Q4 FY28 > 24 months", "Investment payback vs FY26-flat > 30 months (FY29)"]:
    T.cell(r, 1, t).font = BLK; T.cell(r, 1).alignment = Alignment(wrap_text=True)
    for j, dl in enumerate(["Win rate (pts)", "Churn (pts)", "Plan delivery (share)"]):
        v = be.get(f"{t} | {dl}", ""); c = T.cell(r, 2 + j, v); c.font = BOLD
        if isinstance(v, float): c.number_format = "0%" if "delivery" in dl else "+0.0;−0.0"
    T.row_dimensions[r].height = 28; r += 1
T.cell(r + 1, 1, "Reading: the every-quarter payback test is the binding one. Q3 FY27 sits at 29.5 months, so a win rate half a point below plan breaks it; hence the Q2 FY27 trigger that holds the H2 reserve if win rate is off track. Sales-cycle gains add little because pipeline supply, not AE capacity, binds from Q3 FY27. Churn does not move CAC payback (it is measured on new and expansion ARR) but it drives NRR and FY29 ARR more than any other driver.").font = SM
T.cell(r + 1, 1).alignment = Alignment(wrap_text=True, vertical="top"); T.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=10); T.row_dimensions[r + 1].height = 60
# Cascade saved values
X = wb["Cascade"]
for k, (b, p) in enumerate(zip(runs["casc_base"], plus3)):
    X[f"C{5 + k}"] = b; X[f"D{5 + k}"] = p
    X[f"C{5 + k}"].font = BLK; X[f"D{5 + k}"].font = BLK
# Traceability
Z = wb.create_sheet("Traceability"); Z.column_dimensions["A"].width = 40; Z.column_dimensions["B"].width = 30; Z.column_dimensions["C"].width = 44; Z.column_dimensions["D"].width = 20
Z["A1"] = "Traceability map: every major assumption and where it came from"; Z["A1"].font = TT
hdr(Z, 3, ["Assumption (model location)", "Value used (Base)", "Evidence and source", "Owner deliverable"])
TR = [("Opening ARR and FY26 bridge (ARR_bridge B12)", "₹310 cr; +41 new, +22 exp, −9 con, −16 churn", "DR3 pl_quarterly; DR2 arr_movements", "D1"),
 ("Direct win rate FY26 (Assumptions)", "17.1%", "DR1 cleaned (30 duplicates, 16 tests, 88 DQ-as-lost removed; D1 rules)", "D1"),
 ("Win-rate path", "19.5% / 21.5% / 22.5%", "D2 tier mix (+2.9 pts at full effect); D5 battlecards (+0.6 pts beyond D2); Inject 2 holds Helix-incumbent win at 7%", "D2, D5"),
 ("Sales cycle", "8.1 → 7.3 months", "DR1 stage history; D4 stage-exit rules; D3 business-case discipline", "D1, D3, D4"),
 ("Deal size and discount", "₹71.4 → 77 lakh list; 19.3% → 15%", "DR1 FY26 won deals; D3 negotiation plan; D5 discount limits (cap 15%)", "D3, D5"),
 ("Win-rate response to discount", "0.1 pt per pt", "D1 discount-creep test; D5 DISCOUNT-NO-EFFECT interview code (to validate in interviews)", "D1, D5"),
 ("Qualified pipeline supply", "380 / 430 / 470 deals", "DR1 FY26 392; D2 tiering and ₹24 cr programme reallocation", "D2"),
 ("Deals an AE carries at once", "5.4", "DR3 ae_quarterly (48.5 ramped-AE years) × DR1 closes and cycle", "D1, D2"),
 ("Ramp weights by tenure", "0 / 0.33 / 0.66 / 1.0", "D1 productivity by tenure ₹0 / 0.29 / 0.59 / 0.89 cr (DR3); management's 6-month ramp rejected", "D1"),
 ("AE attrition", "29% → 24% / 22% / 20%", "DR3 roster; Inject 1 AE retention pool", "D1, Inject 1"),
 ("Net new AE hires", "6 (Q3 FY27), 8 (FY28), 6 (FY29)", "Inject 1 re-prioritisation; guardrail ≤ 12 in FY27", "Inject 1, D2"),
 ("Q1 FY27 new-logo ARR", "₹8.71 cr", "D4 locked hold-out forecast (sha256 ba5b63ad…)", "D4"),
 ("Partner activations, productivity, ramp", "8 / 9 / 5; ₹1.2 cr; 15% / 60% / 100%", "DR6; D1 SEA partner win 34%; D6 economics", "D6"),
 ("Existing partner ARR after Inject 3", "₹2.26 cr", "DR1 open pipeline; P05/P11 haircut", "D6, Inject 3"),
 ("Partner margins and channel cost", "16–18% year 1; ₹2.44 / 3.46 / 4.0 cr fixed", "D6 programme design; DR3 partner manager cost", "D6"),
 ("Expansion, contraction, churn", "8.5→10.5%; 3.0→2.5%; 5.0→4.0%", "DR2/DR3 movements; D1 NRR diagnosis (Helix NRR 116% → 92%); Inject 1 retention squad; Inject 2", "D1, Inject 1"),
 ("Vajra re-tender", "20% loss probability; 8% concession", "Inject 4; DR2 CU-1095; D5 Helix battlecard", "D7 (Inject 4)"),
 ("Hardware attach", "1.48 → 1.05 per ₹1 new + exp ARR", "DR3 FY24–26 ratios; D1 SENSOR-REUSE interview code; sensor-agnostic tier from Q4 FY27", "D1, D8"),
 ("Cost of revenue", "Software 18%; hardware 75%", "DR3 FY26; DR3 unit costs", "DR3"),
 ("S&M base, AE cost, recruiting", "₹60.3 cr other lines; ₹0.56 cr per AE-year; ₹3.5 lakh per hire", "DR3 sm_by_line; unit costs", "DR3"),
 ("Investment lines FY27", "₹14.95 cr of ₹15 cr", "Inject 1 envelope workbook (reserve released at Q2 review)", "Inject 1"),
 ("R&D and G&A growth", "6% / 4% a year", "Assumption; no new hardware (guardrail)", "D7"),
 ("Do-nothing counterfactual", "New logo −5% a year; NRR −1 pt a year", "D1 trends FY24–26 (win 27% → 17%; NRR 118% → 99%)", "D1"),
 ("Guardrails and board path", "15/20/22%; GM 66%; payback 30/24; S&M 30%", "Brief §4 and §6", "Brief"),
 ("CAC payback definition", "S&M ÷ ((new + exp ARR) × GM) × 12", "DR3 definitions (board-pack method)", "DR3"),
 ("Primary evidence", "8 interviews (in progress)", "Interview kit and log; codes cited above are flagged until interviews are coded", "Kit")]
for i, row in enumerate(TR):
    for j, v in enumerate(row):
        c = Z.cell(4 + i, 1 + j, v); c.font = BLK; c.alignment = Alignment(wrap_text=True, vertical="top")
order = ["README", "Assumptions", "Scenarios", "Sensitivity", "Cascade", "Checks", "Capacity", "Bookings", "Channel", "ARR_bridge", "PnL", "Investment", "Do_nothing", "Inject4", "Traceability"]
wb._sheets = [wb[n] for n in order]
wb.save(F); print("finalized")
